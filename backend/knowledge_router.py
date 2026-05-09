"""RAG / Knowledge Base router."""
from fastapi import APIRouter, HTTPException, Depends, UploadFile, File
from pydantic import BaseModel, Field
from typing import List, Optional, Dict, Any
import io
import re

from core import db, get_current_user, User, embed_batch, embed_text, tokenize, bm25_score, now_iso, new_id, logger

import httpx
from pypdf import PdfReader
import docx as docx_lib
from bs4 import BeautifulSoup
from collections import Counter
import math


router = APIRouter(prefix="/knowledge", tags=["knowledge"])


# ======== Models ========

class KnowledgeBase(BaseModel):
    id: str = Field(default_factory=lambda: new_id())
    user_id: str
    name: str
    description: str = ""
    doc_count: int = 0
    chunk_count: int = 0
    created_at: str = Field(default_factory=now_iso)
    updated_at: str = Field(default_factory=now_iso)


class Document(BaseModel):
    id: str = Field(default_factory=lambda: new_id())
    kb_id: str
    user_id: str
    name: str
    source_type: str  # 'pdf' | 'docx' | 'txt' | 'url'
    source: str = ""  # url or filename
    char_count: int = 0
    chunk_count: int = 0
    created_at: str = Field(default_factory=now_iso)


class CreateKBRequest(BaseModel):
    name: str
    description: str = ""


class IngestUrlRequest(BaseModel):
    url: str


class SearchRequest(BaseModel):
    query: str
    top_k: int = 5


# ======== Helpers ========

CHUNK_SIZE = 800
CHUNK_OVERLAP = 100


def chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> List[str]:
    text = re.sub(r"\s+", " ", text).strip()
    if len(text) <= size:
        return [text] if text else []
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        chunk = text[start:end]
        # try to break on sentence boundary
        if end < len(text):
            last_period = chunk.rfind(". ")
            if last_period > size * 0.5:
                chunk = chunk[:last_period + 1]
                end = start + last_period + 1
        chunks.append(chunk.strip())
        start = end - overlap
        if start <= 0 or start >= len(text):
            break
    return [c for c in chunks if c]


def extract_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n\n".join(parts)


def extract_docx(content: bytes) -> str:
    doc = docx_lib.Document(io.BytesIO(content))
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_txt(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("latin-1", errors="ignore")


async def fetch_url_text(url: str) -> str:
    async with httpx.AsyncClient(follow_redirects=True, timeout=20) as client:
        r = await client.get(url, headers={"User-Agent": "Mozilla/5.0 FORGE-RAG/1.0"})
        r.raise_for_status()
        soup = BeautifulSoup(r.text, "html.parser")
        for tag in soup(["script", "style", "nav", "footer", "iframe", "noscript"]):
            tag.decompose()
        text = soup.get_text(separator="\n")
        text = re.sub(r"\n\s*\n", "\n\n", text)
        return text.strip()


async def index_document(doc: Document, full_text: str) -> int:
    chunks = chunk_text(full_text)
    if not chunks:
        return 0
    docs_to_insert = []
    for i, chunk in enumerate(chunks):
        toks = tokenize(chunk)
        docs_to_insert.append({
            "id": new_id(),
            "kb_id": doc.kb_id,
            "doc_id": doc.id,
            "doc_name": doc.name,
            "user_id": doc.user_id,
            "position": i,
            "content": chunk,
            "tokens": toks,
            "created_at": now_iso(),
        })
    if docs_to_insert:
        await db.kb_chunks.insert_many(docs_to_insert)
    return len(docs_to_insert)


async def search_chunks(kb_id: str, user_id: str, query: str, top_k: int = 5) -> List[Dict[str, Any]]:
    """BM25 keyword search over chunks."""
    kb = await db.knowledge_bases.find_one({"id": kb_id, "user_id": user_id}, {"_id": 0})
    if not kb:
        return []
    q_tokens = tokenize(query)
    if not q_tokens:
        return []
    cursor = db.kb_chunks.find({"kb_id": kb_id, "user_id": user_id}, {"_id": 0})
    chunks = await cursor.to_list(10000)
    if not chunks:
        return []

    # Compute idf
    N = len(chunks)
    df = Counter()
    total_len = 0
    for c in chunks:
        toks = c.get("tokens") or tokenize(c.get("content", ""))
        total_len += len(toks)
        for t in set(toks):
            df[t] += 1
    avgdl = total_len / max(N, 1)
    idf = {t: math.log(1 + (N - df_t + 0.5) / (df_t + 0.5)) for t, df_t in df.items()}

    scored = []
    for c in chunks:
        toks = c.get("tokens") or tokenize(c.get("content", ""))
        score = bm25_score(q_tokens, toks, idf, avgdl)
        if score > 0:
            scored.append({"score": score, **c})
    scored.sort(key=lambda x: x["score"], reverse=True)
    top = scored[:top_k]
    for c in top:
        c.pop("tokens", None)
    return top


# ======== Routes ========

@router.get("", response_model=List[KnowledgeBase])
async def list_knowledge_bases(user: User = Depends(get_current_user)):
    docs = await db.knowledge_bases.find({"user_id": user.user_id}, {"_id": 0}).sort("updated_at", -1).to_list(200)
    return [KnowledgeBase(**d) for d in docs]


@router.post("", response_model=KnowledgeBase)
async def create_kb(req: CreateKBRequest, user: User = Depends(get_current_user)):
    kb = KnowledgeBase(user_id=user.user_id, name=req.name, description=req.description)
    await db.knowledge_bases.insert_one(kb.model_dump())
    return kb


@router.get("/{kb_id}", response_model=KnowledgeBase)
async def get_kb(kb_id: str, user: User = Depends(get_current_user)):
    doc = await db.knowledge_bases.find_one({"id": kb_id, "user_id": user.user_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Knowledge base not found")
    return KnowledgeBase(**doc)


@router.delete("/{kb_id}")
async def delete_kb(kb_id: str, user: User = Depends(get_current_user)):
    result = await db.knowledge_bases.delete_one({"id": kb_id, "user_id": user.user_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Knowledge base not found")
    await db.kb_documents.delete_many({"kb_id": kb_id})
    await db.kb_chunks.delete_many({"kb_id": kb_id})
    return {"success": True}


@router.get("/{kb_id}/documents", response_model=List[Document])
async def list_documents(kb_id: str, user: User = Depends(get_current_user)):
    kb = await db.knowledge_bases.find_one({"id": kb_id, "user_id": user.user_id}, {"_id": 0})
    if not kb:
        raise HTTPException(status_code=404, detail="Knowledge base not found")
    docs = await db.kb_documents.find({"kb_id": kb_id}, {"_id": 0}).sort("created_at", -1).to_list(500)
    return [Document(**d) for d in docs]


async def _refresh_kb_counts(kb_id: str):
    doc_count = await db.kb_documents.count_documents({"kb_id": kb_id})
    chunk_count = await db.kb_chunks.count_documents({"kb_id": kb_id})
    await db.knowledge_bases.update_one(
        {"id": kb_id},
        {"$set": {"doc_count": doc_count, "chunk_count": chunk_count, "updated_at": now_iso()}},
    )


@router.post("/{kb_id}/upload", response_model=Document)
async def upload_document(kb_id: str, file: UploadFile = File(...), user: User = Depends(get_current_user)):
    kb = await db.knowledge_bases.find_one({"id": kb_id, "user_id": user.user_id}, {"_id": 0})
    if not kb:
        raise HTTPException(status_code=404, detail="Knowledge base not found")

    content = await file.read()
    if len(content) > 15 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File too large (max 15MB)")

    filename = file.filename or "upload"
    lower = filename.lower()
    try:
        if lower.endswith(".pdf"):
            text, source_type = extract_pdf(content), "pdf"
        elif lower.endswith(".docx"):
            text, source_type = extract_docx(content), "docx"
        elif lower.endswith((".txt", ".md")):
            text, source_type = extract_txt(content), "txt"
        else:
            raise HTTPException(status_code=400, detail="Unsupported file type. Use PDF, DOCX, TXT or MD.")
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("Extract failed")
        raise HTTPException(status_code=400, detail=f"Extraction failed: {e}")

    if not text or len(text.strip()) < 10:
        raise HTTPException(status_code=400, detail="No extractable text found")

    doc = Document(
        kb_id=kb_id, user_id=user.user_id, name=filename,
        source_type=source_type, source=filename, char_count=len(text),
    )
    await db.kb_documents.insert_one(doc.model_dump())

    try:
        chunk_count = await index_document(doc, text)
        doc.chunk_count = chunk_count
        await db.kb_documents.update_one({"id": doc.id}, {"$set": {"chunk_count": chunk_count}})
    except Exception as e:
        logger.exception("Index failed")
        await db.kb_documents.delete_one({"id": doc.id})
        raise HTTPException(status_code=500, detail=f"Indexing failed: {e}")

    await _refresh_kb_counts(kb_id)
    return doc


@router.post("/{kb_id}/url", response_model=Document)
async def ingest_url(kb_id: str, req: IngestUrlRequest, user: User = Depends(get_current_user)):
    kb = await db.knowledge_bases.find_one({"id": kb_id, "user_id": user.user_id}, {"_id": 0})
    if not kb:
        raise HTTPException(status_code=404, detail="Knowledge base not found")

    try:
        text = await fetch_url_text(req.url)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"URL fetch failed: {e}")

    if not text or len(text.strip()) < 50:
        raise HTTPException(status_code=400, detail="No extractable text from URL")

    name = req.url
    doc = Document(
        kb_id=kb_id, user_id=user.user_id, name=name,
        source_type="url", source=req.url, char_count=len(text),
    )
    await db.kb_documents.insert_one(doc.model_dump())

    try:
        chunk_count = await index_document(doc, text)
        doc.chunk_count = chunk_count
        await db.kb_documents.update_one({"id": doc.id}, {"$set": {"chunk_count": chunk_count}})
    except Exception as e:
        logger.exception("Index failed")
        await db.kb_documents.delete_one({"id": doc.id})
        raise HTTPException(status_code=500, detail=f"Indexing failed: {e}")

    await _refresh_kb_counts(kb_id)
    return doc


@router.delete("/{kb_id}/documents/{doc_id}")
async def delete_document(kb_id: str, doc_id: str, user: User = Depends(get_current_user)):
    result = await db.kb_documents.delete_one({"id": doc_id, "kb_id": kb_id, "user_id": user.user_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Document not found")
    await db.kb_chunks.delete_many({"doc_id": doc_id})
    await _refresh_kb_counts(kb_id)
    return {"success": True}


@router.post("/{kb_id}/search")
async def search_kb(kb_id: str, req: SearchRequest, user: User = Depends(get_current_user)):
    chunks = await search_chunks(kb_id, user.user_id, req.query, top_k=req.top_k)
    return {"results": chunks}
